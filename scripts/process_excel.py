"""
Process organized_data.xlsx and export structured JSON for SkillTrace frontend.
"""
import openpyxl
import json
import os
from datetime import datetime

EXCEL_PATH = os.path.abspath('organized_data.xlsx')
BANK_PATH = os.path.abspath('bank_statements.xlsx')
OUTPUT_DIR = os.path.abspath('frontend/src/api/mock/data')
OUTPUT_PATH = os.path.join(OUTPUT_DIR, 'organized_data.json')

AS_OF = '2026-09-10'

# 36 months from 2023-10 to 2026-09
BANK_MONTHS = []
_y, _m = 2023, 10
for _ in range(36):
    BANK_MONTHS.append(f"{_y}-{str(_m).zfill(2)}")
    _m += 1
    if _m > 12:
        _m = 1
        _y += 1

def get_bank_month_index(date_str):
    if not date_str:
        return None
    try:
        d = datetime.strptime(date_str[:10], '%Y-%m-%d')
        idx = (d.year - 2023) * 12 + (d.month - 10)
        return idx if 0 <= idx < 36 else None
    except Exception:
        return None

def days_between(a_str, b_str):
    d_a = datetime.strptime(a_str, '%Y-%m-%d')
    d_b = datetime.strptime(b_str, '%Y-%m-%d')
    return (d_b - d_a).days

def trust_for(source, date_str):
    if days_between(date_str, AS_OF) > 275:
        return 'stale'
    if source in ('bank', 'employer'):
        return 'high'
    if source == 'field_officer':
        return 'medium'
    return 'low'

DRIFT_ROLES = [
    'Helper / Unskilled', 'Delivery Partner', 'Security Guard', 'Retail Cashier',
    'Machine Operator', 'Site Supervisor', 'Office Assistant', 'Driver',
]

COURSE_DEFS = {
    'AI & Machine Learning': {'intended_role': 'AI / ML Engineer', 'target_pct': 85},
    'Accounting & GST': {'intended_role': 'Accounts Executive', 'target_pct': 80},
    'Automobile Technician': {'intended_role': 'Automobile Technician', 'target_pct': 78},
    'Beauty & Wellness': {'intended_role': 'Beauty Therapist', 'target_pct': 75},
    'CNC Operator': {'intended_role': 'CNC Operator', 'target_pct': 82},
    'Cloud Computing': {'intended_role': 'Cloud Support Associate', 'target_pct': 80},
    'Cybersecurity': {'intended_role': 'Cybersecurity Analyst', 'target_pct': 82},
    'Data Analytics': {'intended_role': 'Data Analyst', 'target_pct': 84},
    'Digital Marketing': {'intended_role': 'Digital Marketing Executive', 'target_pct': 76},
    'Electrician': {'intended_role': 'Electrician', 'target_pct': 80},
    'Entrepreneurship': {'intended_role': 'Business Owner', 'target_pct': 70},
    'Fitter': {'intended_role': 'Mechanical Fitter', 'target_pct': 80},
    'Graphic Design': {'intended_role': 'Graphic Designer', 'target_pct': 74},
    'Healthcare Assistant': {'intended_role': 'Healthcare Assistant', 'target_pct': 75},
    'Logistics & Supply Chain': {'intended_role': 'Logistics Executive', 'target_pct': 78},
    'Mobile App Development': {'intended_role': 'Mobile App Developer', 'target_pct': 82},
    'Python Development': {'intended_role': 'Python Developer', 'target_pct': 85},
    'Retail Sales': {'intended_role': 'Sales Associate', 'target_pct': 75},
    'Solar Technician': {'intended_role': 'Solar Technician', 'target_pct': 72},
    'Web Development': {'intended_role': 'Web Developer', 'target_pct': 82},
}

CATEGORIES = ['General', 'OBC', 'SC', 'ST']

PROVIDER_DISTRICT_MAP = {
    'Digital India Academy': 'Thane',
    'Future Skills Institute': 'Nagpur',
    'Gujarat Skill Development Centre': 'Pune',
    'Pragati Skill Centre': 'Satara',
    'Skill India Training Centre': 'Mumbai Suburban',
    'TechSkills Academy': 'Nashik',
    'Udaan Training Foundation': 'Nanded',
}

def main():
    print(f"Reading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    ws_t = wb['trainees']
    ws_eo = wb['employment_outcomes']
    ws_f = wb['followups']
    ws_a = wb['assessments']
    ws_te = wb['training_enrollments']
    ws_emp = wb['employers']
    
    # 1. Employers map
    employers_map = {}
    for r in ws_emp.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        employers_map[r[0]] = {
            'id': r[0],
            'company_name': r[1],
            'industry': r[2],
            'district': r[3],
            'company_size': r[4],
            'verified': r[5] == 'Yes'
        }
    print(f"Loaded {len(employers_map)} employers.")

    # 2. Training enrollments map
    enrollments_map = {}
    for r in ws_te.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        t_id = r[0].replace('TRG', 'TRN')
        start_date = r[3].strftime('%Y-%m-%d') if r[3] else None
        comp_date = r[4].strftime('%Y-%m-%d') if r[4] else None
        q = None
        if r[4]:
            q_num = (r[4].month - 1) // 3 + 1
            q = f"{r[4].year}-Q{q_num}"
        enrollments_map[t_id] = {
            'start_date': start_date,
            'completion_date': comp_date,
            'attendance_percentage': r[5],
            'completion_status': r[6],
            'cohort': q or '2025-Q4'
        }
    print(f"Loaded {len(enrollments_map)} training enrollments.")

    # 3. Assessments map
    assessments_map = {}
    for r in ws_a.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        assessments_map[r[1]] = {
            'technical_score': r[2],
            'soft_skill_score': r[3],
            'assessment_result': r[4],
            'skill_level': r[5],
            'certified': r[6] == 'Yes'
        }
    print(f"Loaded {len(assessments_map)} assessments.")

    # 4. Employment outcomes map
    outcomes_map = {}
    for r in ws_eo.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        start_date = r[4].strftime('%Y-%m-%d') if r[4] else None
        outcomes_map[r[1]] = {
            'outcome_id': r[0],
            'status': r[2],
            'employer_id': r[3],
            'start_date': start_date,
            'initial_income': r[5],
            'current_income': r[6],
            'retention_months': r[7],
            'job_relevant': r[8],
            'currently_employed': r[9] == 'Yes'
        }
    print(f"Loaded {len(outcomes_map)} employment outcomes.")

    # 5. Followups map
    followups_map = {}
    for r in ws_f.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        t_id = r[1]
        f_date = r[2].strftime('%Y-%m-%d') if r[2] else None
        item = {
            'followup_id': r[0],
            'followup_date': f_date,
            'months_after_training': r[3],
            'contacted': r[4] == 'Yes',
            'current_status': r[5],
            'monthly_income': r[6],
            'job_satisfaction': r[7],
            'training_relevance': r[8],
            'skill_gap_identified': r[9],
            'reason_for_attrition': r[10]
        }
        followups_map.setdefault(t_id, []).append(item)
    for t_id in followups_map:
        followups_map[t_id].sort(key=lambda x: x['followup_date'] or '')
    print(f"Loaded followups for {len(followups_map)} trainees.")

    # 6. Providers
    provider_names = sorted(list(set(r[8] for r in ws_t.iter_rows(min_row=2, values_only=True) if r[8])))
    provider_to_id = {name: f"PRV-{str(i+1).zfill(3)}" for i, name in enumerate(provider_names)}
    providers_list = [
        {
            'id': provider_to_id[name],
            'name': name,
            'district': PROVIDER_DISTRICT_MAP.get(name, 'Pune')
        }
        for name in provider_names
    ]

    # 6b. Bank statements map
    print(f"Reading {BANK_PATH}...")
    wb_bank = openpyxl.load_workbook(BANK_PATH, data_only=True, read_only=True)
    ws_b = wb_bank['Monthly Income']
    bank_rows = list(ws_b.iter_rows(values_only=True))
    bank_headers = bank_rows[0]
    bank_map = {}
    for col_idx, tid in enumerate(bank_headers):
        if not tid:
            continue
        raw_vals = [bank_rows[r][col_idx] for r in range(1, len(bank_rows))]
        cleaned_series = []
        for v in raw_vals:
            if isinstance(v, (int, float)):
                cleaned_series.append(int(v))
            else:
                cleaned_series.append(None)
        
        pos_vals = [v for v in cleaned_series if v is not None and v > 0]
        months_credited = len(pos_vals)
        avg_income = round(sum(pos_vals) / months_credited) if months_credited else 0
        latest_income = 0
        for v in reversed(cleaned_series):
            if v is not None and v > 0:
                latest_income = v
                break

        is_all_none = all(v is None for v in cleaned_series)
        if is_all_none:
            status = 'unbanked'
        elif months_credited >= 6:
            status = 'retained'
        elif months_credited >= 3:
            status = 'active'
        elif months_credited > 0:
            status = 'irregular'
        else:
            status = 'inactive'

        bank_map[tid] = {
            'summary': {
                'verified': months_credited >= 3,
                'months_credited': months_credited,
                'avg_income': avg_income,
                'latest_income': latest_income,
                'status': status
            },
            'series': cleaned_series
        }
    print(f"Loaded bank statements for {len(bank_map)} trainees.")

    # 7. Build Trainees and Events
    trainees = []
    events = []
    event_seq = 1

    districts_set = set()
    cohorts_set = set()

    for idx, r in enumerate(ws_t.iter_rows(min_row=2, values_only=True)):
        if not r[0]: continue
        t_id = r[0]
        name = r[1]
        age = r[2]
        gender = r[3]
        district = r[4]
        education = r[5]
        phone_num = r[6]
        email = r[7]
        prov_name = r[8]
        prog_name = r[9]
        enrol_date = r[10].strftime('%Y-%m-%d') if r[10] else None

        districts_set.add(district)

        # Age group
        if age <= 24:
            age_group = '18-24'
        elif age <= 34:
            age_group = '25-34'
        else:
            age_group = '35-44'

        # Category deterministic distribution
        cat_hash = (int(t_id.replace('TRN', '')) + 7) % 100
        if cat_hash < 30:
            category = 'General'
        elif cat_hash < 66:
            category = 'OBC'
        elif cat_hash < 88:
            category = 'SC'
        else:
            category = 'ST'

        enr = enrollments_map.get(t_id, {})
        cohort = enr.get('cohort', '2025-Q3')
        cohorts_set.add(cohort)
        comp_date = enr.get('completion_date') or enrol_date or '2025-06-30'

        prov_id = provider_to_id.get(prov_name, 'PRV-001')
        phone = f"+91 {str(phone_num)[:5]} {str(phone_num)[5:]}" if phone_num else "+91 98000 00000"

        c_def = COURSE_DEFS.get(prog_name, {'intended_role': prog_name, 'target_pct': 75})
        intended_role = c_def['intended_role']

        ass = assessments_map.get(t_id, {})
        certified = ass.get('certified', True)

        b_info = bank_map.get(t_id, {
            'summary': {
                'verified': False,
                'months_credited': 0,
                'avg_income': 0,
                'latest_income': 0,
                'status': 'unbanked'
            },
            'series': [None] * 36
        })

        trainees.append({
            'id': t_id,
            'name': name,
            'course': prog_name,
            'district': district,
            'gender': gender,
            'age_group': age_group,
            'category': category,
            'cohort': cohort,
            'provider_id': prov_id,
            'phone': phone,
            'education': education,
            'email': email,
            'certified': certified,
            'bank_summary': b_info['summary'],
            'bank_series': b_info['series']
        })

        # Build events for this trainee
        eo = outcomes_map.get(t_id, {})
        fups = followups_map.get(t_id, [])
        emp_obj = employers_map.get(eo.get('employer_id'), {})
        emp_name = emp_obj.get('company_name') if emp_obj else None
        is_verified_emp = emp_obj.get('verified', False)

        on_role = (eo.get('job_relevant') == 'Yes')
        job_role = intended_role if on_role else DRIFT_ROLES[idx % len(DRIFT_ROLES)]

        placed_date = eo.get('start_date')
        if not placed_date and eo.get('status') == 'Employed':
            placed_date = comp_date

        placed_recorded = False
        if placed_date:
            placed_m_idx = get_bank_month_index(placed_date)
            bank_credit_at_placement = (
                b_info['series'][placed_m_idx]
                if placed_m_idx is not None and b_info['series'][placed_m_idx] is not None and b_info['series'][placed_m_idx] > 0
                else None
            )

            if is_verified_emp:
                source = 'employer'
            elif bank_credit_at_placement:
                source = 'bank'
            else:
                source = 'trainee'

            salary = bank_credit_at_placement or eo.get('initial_income') or 12000
            events.append({
                'id': f"EVT-{str(event_seq).zfill(5)}",
                'trainee_id': t_id,
                'date': placed_date,
                'what_happened': 'placed',
                'job_role': job_role,
                'salary': salary,
                'source': source,
                'trust_level': trust_for(source, placed_date),
                'employer': emp_name or 'Maharashtra Industrial Solutions'
            })
            event_seq += 1
            placed_recorded = True

        # Process followups into events
        for f in fups:
            f_date = f['followup_date']
            if not f_date or not f['contacted']:
                continue
            
            f_m_idx = get_bank_month_index(f_date)
            bank_credit = (
                b_info['series'][f_m_idx]
                if f_m_idx is not None and b_info['series'][f_m_idx] is not None and b_info['series'][f_m_idx] > 0
                else None
            )

            f_status = f['current_status']
            if f_status == 'Employed':
                # If previously placed, this is still_working
                if placed_recorded:
                    if is_verified_emp:
                        source = 'employer'
                    elif bank_credit:
                        source = 'bank'
                    else:
                        source = 'trainee'

                    salary = bank_credit or f['monthly_income'] or eo.get('current_income') or 15000
                    events.append({
                        'id': f"EVT-{str(event_seq).zfill(5)}",
                        'trainee_id': t_id,
                        'date': f_date,
                        'what_happened': 'still_working',
                        'job_role': job_role,
                        'salary': salary,
                        'source': source,
                        'trust_level': trust_for(source, f_date),
                        'employer': emp_name or 'Maharashtra Industrial Solutions'
                    })
                    event_seq += 1
            elif f_status == 'Self-employed':
                if bank_credit:
                    source = 'bank'
                else:
                    source = 'trainee'

                salary = bank_credit or f['monthly_income'] or 12000
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': f_date,
                    'what_happened': 'self_employed',
                    'job_role': f"Self-employed — {intended_role}",
                    'salary': salary,
                    'source': source,
                    'trust_level': trust_for(source, f_date),
                    'employer': None
                })
                event_seq += 1
            elif f_status == 'Apprentice':
                if is_verified_emp:
                    source = 'employer'
                elif bank_credit:
                    source = 'bank'
                else:
                    source = 'trainee'

                salary = bank_credit or f['monthly_income'] or 8000
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': f_date,
                    'what_happened': 'apprentice',
                    'job_role': f"Apprentice — {intended_role}",
                    'salary': salary,
                    'source': source,
                    'trust_level': trust_for(source, f_date),
                    'employer': emp_name
                })
                event_seq += 1
            elif f_status in ('Unemployed', 'Dropped out'):
                what = 'left_job' if placed_recorded else 'not_working'
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': f_date,
                    'what_happened': what,
                    'job_role': None,
                    'salary': None,
                    'source': 'trainee',
                    'trust_level': trust_for('trainee', f_date),
                    'employer': emp_name if what == 'left_job' else None
                })
                event_seq += 1

        # If no placement and no followup event, add initial status from outcome
        trainee_evts = [e for e in events if e['trainee_id'] == t_id]
        if not trainee_evts:
            comp_m_idx = get_bank_month_index(comp_date)
            bank_credit = (
                b_info['series'][comp_m_idx]
                if comp_m_idx is not None and b_info['series'][comp_m_idx] is not None and b_info['series'][comp_m_idx] > 0
                else None
            )

            eo_status = eo.get('status')
            if eo_status == 'Self-employed':
                source = 'bank' if bank_credit else 'trainee'
                salary = bank_credit or eo.get('initial_income') or 11000
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': comp_date,
                    'what_happened': 'self_employed',
                    'job_role': f"Self-employed — {intended_role}",
                    'salary': salary,
                    'source': source,
                    'trust_level': trust_for(source, comp_date),
                    'employer': None
                })
                event_seq += 1
            elif eo_status == 'Apprentice':
                source = 'employer' if is_verified_emp else ('bank' if bank_credit else 'trainee')
                salary = bank_credit or eo.get('initial_income') or 7500
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': comp_date,
                    'what_happened': 'apprentice',
                    'job_role': f"Apprentice — {intended_role}",
                    'salary': salary,
                    'source': source,
                    'trust_level': trust_for(source, comp_date),
                    'employer': emp_name
                })
                event_seq += 1
            elif eo_status in ('Unemployed', 'Dropped out'):
                events.append({
                    'id': f"EVT-{str(event_seq).zfill(5)}",
                    'trainee_id': t_id,
                    'date': comp_date,
                    'what_happened': 'not_working',
                    'job_role': None,
                    'salary': None,
                    'source': 'trainee',
                    'trust_level': trust_for('trainee', comp_date),
                    'employer': None
                })
                event_seq += 1

    print(f"Built {len(trainees)} trainees and {len(events)} events.")

    # 8. Disputes
    candidates = [t for t in trainees if any(e['what_happened'] == 'placed' for e in events if e['trainee_id'] == t['id'])]
    disputes = []
    CLAIM_PAIRS = [
        {
            'employer_claim': 'Left employment on 12 Mar 2026 — did not complete notice period.',
            'trainee_claim': 'Still working at the same unit; salary credited in March and April.',
        },
        {
            'employer_claim': 'Never joined after offer letter was issued.',
            'trainee_claim': 'Joined on 04 Jan 2026, worked 5 weeks, paid in cash without payslip.',
        },
        {
            'employer_claim': 'Employed as Helper, monthly wage ₹9,500.',
            'trainee_claim': 'Working as specialist technician, monthly wage ₹16,000.',
        },
        {
            'employer_claim': 'Contract ended at 2 months; not renewed.',
            'trainee_claim': 'Contract renewed verbally; continued for 7 months.',
        },
        {
            'employer_claim': 'Absent without notice since 20 Feb 2026.',
            'trainee_claim': 'On approved medical leave; rejoined 05 Mar 2026.',
        },
        {
            'employer_claim': 'Working part-time, 4 days a week.',
            'trainee_claim': 'Full-time, 6 days a week including Saturdays.',
        },
        {
            'employer_claim': 'Trainee resigned voluntarily in Dec 2025.',
            'trainee_claim': 'Was told not to return after the unit reduced its workforce.',
        },
        {
            'employer_claim': 'Wage band ₹10,000–₹15,000, paid by bank transfer.',
            'trainee_claim': 'Paid ₹9,000 in cash; no bank transfer received.',
        },
        {
            'employer_claim': 'Apprentice, stipend only — not a placement.',
            'trainee_claim': 'Full employee doing the same work as permanent staff.',
        },
    ]

    for i in range(12):
        t = candidates[(i * 37) % len(candidates)]
        pair = CLAIM_PAIRS[i % len(CLAIM_PAIRS)]
        t_evts = [e for e in events if e['trainee_id'] == t['id']]
        pl = next((e for e in t_evts if e['what_happened'] == 'placed'), None)
        d_day = 10 + (i * 5) % 60
        disputes.append({
            'id': f"DSP-{str(i+1).zfill(4)}",
            'trainee_id': t['id'],
            'trainee_name': t['name'],
            'course': t['course'],
            'district': t['district'],
            'employer': pl['employer'] if pl else 'Industrial Corp',
            'employer_claim': pair['employer_claim'],
            'trainee_claim': pair['trainee_claim'],
            'date': f"2026-07-{str(d_day % 28 + 1).zfill(2)}",
            'status': 'open' if i < 9 else 'resolved',
            'resolution': None if i < 9 else 'employer_stands',
            'resolved_at': None if i < 9 else '2026-08-15',
            'resolved_by': None if i < 9 else 'Reviewer (demo)'
        })

    # 9. Consents
    consents = [
        {
            'id': f"CNS-{str(i+1).zfill(5)}",
            'trainee_id': t['id'],
            'status': 'granted',
            'granted_date': '2025-01-15',
            'withdrawn_date': None,
            'scopes': [
                'Employment status check-ins by SMS or WhatsApp',
                'Confirmation of employment with the named employer',
                'Salary band (not exact salary) from bank-verified records',
                'Use of anonymised outcomes in government skilling reports',
            ]
        }
        for i, t in enumerate(trainees)
    ]

    # 10. Follow-up queue
    uncontacted_candidates = []
    for t in trainees:
        t_fups = followups_map.get(t['id'], [])
        no_contacts = [f for f in t_fups if not f['contacted']]
        if len(no_contacts) >= 1:
            last_f = no_contacts[-1]
            uncontacted_candidates.append({
                'trainee_id': t['id'],
                'name': t['name'],
                'phone': t['phone'],
                'course': t['course'],
                'district': t['district'],
                'provider_id': t['provider_id'],
                'cohort': t['cohort'],
                'attempts': 3 + len(no_contacts),
                'last_contact_date': last_f['followup_date'] or '2026-06-01',
                'channel': 'WhatsApp' if (len(uncontacted_candidates) % 2 == 0) else 'SMS',
                'assigned_to': None
            })
    followup_queue = uncontacted_candidates[:30]

    # 11. Courses list
    courses_list = [
        {'name': name, 'intended_role': info['intended_role'], 'target_pct': info['target_pct']}
        for name, info in COURSE_DEFS.items()
    ]

    # 12. Compile final payload
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    payload = {
        'as_of': AS_OF,
        'bank_months': BANK_MONTHS,
        'districts': sorted(list(districts_set)),
        'courses': courses_list,
        'cohorts': sorted(list(cohorts_set)),
        'providers': providers_list,
        'trainees': trainees,
        'events': events,
        'disputes': disputes,
        'consents': consents,
        'followupQueue': followup_queue
    }

    print(f"Writing output to {OUTPUT_PATH}...")
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False)
    
    js_path = os.path.join(OUTPUT_DIR, 'organized_data.js')
    print(f"Writing JS module to {js_path}...")
    with open(js_path, 'w', encoding='utf-8') as f:
        f.write("export const organizedData = ")
        json.dump(payload, f, ensure_ascii=False)
        f.write(";\nexport default organizedData;\n")
    
    file_size = os.path.getsize(OUTPUT_PATH)
    print(f"Done! File size: {file_size / 1024:.1f} KB")

if __name__ == '__main__':
    main()
